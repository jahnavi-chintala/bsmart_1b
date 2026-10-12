"""Export the config CSVs to an Excel workbook for business review.

Sheets: MAPPING_TABLE (all rows), ENTITY_CONFIG, STUDY_CONFIG, and one RESOLVED_<STUDYID> sheet per
study showing the effective mapping (DEFAULT rows merged with that study's overrides).

    python scripts/export_mapping_xlsx.py [output.xlsx]
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))
from bsmart_gold.config import ConfigSet  # noqa: E402

CONFIG = os.path.join(ROOT, "config")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
WIDE = {"ETL_LOGIC": 90, "FROM_SQL": 80, "WHERE_SQL": 70, "MAPPING_COLUMN_NAME": 30, "COMMENTS": 50}


def add_sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(r)
    for i, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), HEADER_FILL
        ws.column_dimensions[cell.column_letter].width = WIDE.get(name, max(14, len(name) + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main(out):
    cfg = ConfigSet.load(*(os.path.join(CONFIG, f) for f in
                           ("study_config.csv", "entity_config.csv", "mapping_table.csv")))
    wb = Workbook()
    wb.remove(wb.active)

    for title, data in (("MAPPING_TABLE", cfg.mappings), ("ENTITY_CONFIG", cfg.entities),
                        ("STUDY_CONFIG", cfg.studies)):
        header = list(data[0].keys())
        add_sheet(wb, title, header, [[r[h] for h in header] for r in data])

    for studyid in cfg.active_studies():
        header = ["STUDYID", "ENTITY_NAME", "MAPPING_TABLE_NAME", "MAPPING_COLUMN_NAME", "ETL_LOGIC",
                  "TARGET_COLUMN_NAME", "GOLD_TABLE", "DEFINED_AT"]
        rows = [[studyid, p.entity, c.table, c.source_columns, c.logic, c.target, p.gold_table, c.studyid]
                for p in cfg.plan(studyid) for c in p.columns]
        add_sheet(wb, f"RESOLVED_{studyid}"[:31], header, rows)

    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(CONFIG, "B-SMART_Mapping_Table.xlsx"))
